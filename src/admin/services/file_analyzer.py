import hashlib
import re
from io import BytesIO
from typing import Any, Optional

import pandas as pd

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional dependency guard
    load_workbook = None


def infer_column_type(series: pd.Series) -> dict[str, Any]:
    if series.empty or series.count() == 0:
        return {
            "type": "text",
            "pattern": None,
            "sample_values": [],
            "null_count": len(series),
            "unique_count": 0,
        }

    non_null = series.dropna()
    sample = non_null.head(5).tolist()
    null_count = int(series.isna().sum())
    unique_count = int(non_null.nunique()) if len(non_null) else 0

    if pd.api.types.is_numeric_dtype(series):
        if len(non_null) and all(isinstance(v, (int, float)) and 0 <= v <= 1 for v in non_null):
            inferred = "percentage"
        else:
            inferred = "amount"
        return {
            "type": inferred,
            "pattern": None,
            "sample_values": sample,
            "null_count": null_count,
            "unique_count": unique_count,
        }

    if pd.api.types.is_datetime64_any_dtype(series):
        return {
            "type": "date",
            "pattern": "date",
            "sample_values": sample,
            "null_count": null_count,
            "unique_count": unique_count,
        }

    pattern = None
    text_sample = [str(v) for v in non_null.head(10)]
    if text_sample:
        combined = " ".join(text_sample)
        if re.search(r"[A-Z]{5}[0-9]{4}[A-Z]{1}", combined):
            pattern = "pan"
        elif re.search(r"\d{2}[A-Z]{5}\d{4}[A-Z]{1}\d{1}[Z]{1}\d{1}", combined):
            pattern = "gstin"
        elif re.search(r"[A-Z]{4}\d{7}", combined):
            pattern = "utr"
        elif re.search(r"\d{10}", combined):
            pattern = "phone"
        elif re.search(r"[\w\.-]+@[\w\.-]+", combined):
            pattern = "email"
        elif re.search(r"\d{4}-\d{2}-\d{2}", combined) or re.search(r"\d{2}/\d{2}/\d{4}", combined):
            pattern = "date"

    return {
        "type": "text",
        "pattern": pattern,
        "sample_values": sample,
        "null_count": null_count,
        "unique_count": unique_count,
    }


def detect_slab_columns(columns: list[str]) -> dict[str, dict[str, Any]]:
    slab_patterns = [
        r"<.*?₹",
        r"₹.*?–",
        r"₹.*?-",
        r"≥.*?₹",
        r"less than",
        r"greater than",
    ]

    def _tier_bounds(label: str) -> dict[str, float | None]:
        text = str(label or "")
        if "<" in text and "50" in text and ("L" in text or "l" in text):
            return {"tier_min": 0.0, "tier_max": 5000000.0}
        if "50" in text and "1 Cr" in text:
            return {"tier_min": 5000000.0, "tier_max": 10000000.0}
        if "1 Cr" in text and "2 Cr" in text:
            return {"tier_min": 10000000.0, "tier_max": 20000000.0}
        if "2 Cr" in text and "3 Cr" in text:
            return {"tier_min": 20000000.0, "tier_max": 30000000.0}
        if "3 Cr" in text and "5 Cr" in text:
            return {"tier_min": 30000000.0, "tier_max": 50000000.0}
        if "5 Cr" in text and "10 Cr" in text:
            return {"tier_min": 50000000.0, "tier_max": 100000000.0}
        if "10 Cr" in text and "25 Cr" in text:
            return {"tier_min": 100000000.0, "tier_max": 250000000.0}
        if "25 Cr" in text and "50 Cr" in text:
            return {"tier_min": 250000000.0, "tier_max": 500000000.0}
        if "50 Cr" in text and "100 Cr" in text:
            return {"tier_min": 500000000.0, "tier_max": 1000000000.0}
        if ">=" in text or "≥" in text:
            return {"tier_min": 1000000000.0, "tier_max": None}
        return {"tier_min": None, "tier_max": None}

    slabs: dict[str, dict[str, Any]] = {}
    for col in columns:
        if any(re.search(p, col, re.IGNORECASE) for p in slab_patterns):
            bounds = _tier_bounds(col)
            slabs[col] = {
                "is_slab": True,
                "target_table": "RUL_CommissionSlab",
                "target_field": "rate",
                "transformation": "UNPIVOT",
                **bounds,
            }
    return slabs


def _formula_profile(file_bytes: bytes, sheet_name: str, header_row: int, sample_rows: int) -> dict[str, dict[str, Any]]:
    if load_workbook is None:
        return {}

    try:
        workbook_formula = load_workbook(BytesIO(file_bytes), data_only=False, read_only=True)
        workbook_values = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return {}

    if sheet_name not in workbook_formula.sheetnames:
        return {}

    worksheet_formula = workbook_formula[sheet_name]
    worksheet_values = workbook_values[sheet_name]

    max_row = min(int(worksheet_formula.max_row or 0), header_row + sample_rows)
    if max_row <= header_row:
        return {}

    headers: list[str] = []
    for cell in worksheet_formula[header_row]:
        header_value = str(cell.value).strip() if cell.value not in (None, "") else ""
        headers.append(header_value)

    profile: dict[str, dict[str, Any]] = {
        header: {
            "formula_detected": False,
            "formula_count": 0,
            "formula_error_count": 0,
            "formula_samples": [],
            "evaluated_formula_samples": [],
        }
        for header in headers
        if header
    }

    for row_idx in range(header_row + 1, max_row + 1):
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            formula_cell = worksheet_formula.cell(row=row_idx, column=col_idx)
            formula_value = formula_cell.value
            if not (isinstance(formula_value, str) and formula_value.startswith("=")):
                continue

            profile[header]["formula_detected"] = True
            profile[header]["formula_count"] += 1
            if len(profile[header]["formula_samples"]) < 5:
                profile[header]["formula_samples"].append(formula_value)

            evaluated_value = worksheet_values.cell(row=row_idx, column=col_idx).value
            if len(profile[header]["evaluated_formula_samples"]) < 5:
                profile[header]["evaluated_formula_samples"].append(evaluated_value)

            if evaluated_value is None or (isinstance(evaluated_value, str) and evaluated_value.startswith("#")):
                profile[header]["formula_error_count"] += 1

    return profile


def _sheet_analysis(
    df: pd.DataFrame,
    sheet_name: str,
    header_row: int,
    formula_profile: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    fingerprint_data = f"{sheet_name}_{header_row}_{'_'.join(map(str, df.columns.tolist()))}"
    fingerprint = hashlib.sha256(fingerprint_data.encode()).hexdigest()
    formula_profile = formula_profile or {}

    columns: list[dict[str, Any]] = []
    for col in df.columns:
        source_column = str(col)
        formula_data = formula_profile.get(source_column, {})
        col_data = infer_column_type(df[col])
        formula_detected = bool(formula_data.get("formula_detected"))
        formula_error_count = int(formula_data.get("formula_error_count") or 0)
        columns.append(
            {
                "source_column": source_column,
                "inferred_type": col_data["type"],
                "sample_values": col_data["sample_values"],
                "null_count": col_data["null_count"],
                "unique_count": col_data["unique_count"],
                "pattern": col_data.get("pattern"),
                "formula_detected": formula_detected,
                "formula_count": int(formula_data.get("formula_count") or 0),
                "formula_error_count": formula_error_count,
                "formula_samples": formula_data.get("formula_samples", []),
                "evaluated_formula_samples": formula_data.get("evaluated_formula_samples", []),
                "formula_status": "FORMULA_ERRORS" if formula_error_count else ("FORMULA_DETECTED" if formula_detected else "NONE"),
                "suggested_target_table": None,
                "suggested_target_field": None,
                "confidence_score": 0,
                "ignore": False,
                "notes": None,
            }
        )

    return {
        "fingerprint": fingerprint,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "row_count": len(df),
        "columns": columns,
    }


def analyze_excel_bytes(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
    sample_rows: int = 500,
) -> dict[str, Any]:
    xls = pd.ExcelFile(BytesIO(file_bytes))
    requested_sheets = [sheet_name] if sheet_name else list(xls.sheet_names)

    sheets: list[dict[str, Any]] = []
    for sh in requested_sheets:
        df = xls.parse(sheet_name=sh, header=header_row - 1, nrows=sample_rows)
        formula_profile = _formula_profile(file_bytes, sh, header_row, sample_rows)
        sheets.append(_sheet_analysis(df, sh, header_row, formula_profile))

    workbook_fp_data = "|".join([s["fingerprint"] for s in sheets])
    workbook_fingerprint = hashlib.sha256(workbook_fp_data.encode()).hexdigest()

    return {
        "workbook_fingerprint": workbook_fingerprint,
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


def analyze_file(file, sheet_name: Optional[str] = None, header_row: int = 1) -> dict[str, Any]:
    file_bytes = file.file.read()
    return analyze_excel_bytes(file_bytes, sheet_name=sheet_name, header_row=header_row)
