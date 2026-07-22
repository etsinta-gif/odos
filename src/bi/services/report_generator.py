from __future__ import annotations

import csv
import io
from datetime import datetime
from html import escape
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill


class ReportGenerator:
    """Generate reports in HTML, CSV, Excel, and PDF-like output."""

    def __init__(self, data: dict[str, Any], title: str, company_name: str):
        self.data = data
        self.title = title
        self.company_name = company_name
        self.generated_at = datetime.utcnow()

    def to_pdf(self) -> bytes:
        """Generate PDF output. Falls back to plain bytes if reportlab is unavailable."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle("ReportTitle", parent=styles["Heading1"], fontSize=18, spaceAfter=18)

            story = [
                Paragraph(escape(self.title), title_style),
                Paragraph(f"Company: {escape(self.company_name)}", styles["Normal"]),
                Paragraph(f"Generated: {self.generated_at.strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]),
                Spacer(1, 16),
            ]

            columns = self.data.get("columns", [])
            rows = self.data.get("data", [])
            if columns:
                table_data = [columns]
                for row in rows:
                    table_data.append([str(row.get(col, "")) for col in columns])

                table = Table(table_data)
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f4f4f")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ]
                    )
                )
                story.append(table)

            doc.build(story)
            return buffer.getvalue()
        except Exception:
            # Fallback keeps export endpoint functional even without optional PDF deps.
            content = [
                self.title,
                f"Company: {self.company_name}",
                f"Generated: {self.generated_at.strftime('%Y-%m-%d %H:%M:%S')}",
                "",
            ]
            columns = self.data.get("columns", [])
            rows = self.data.get("data", [])
            content.append(",".join(columns))
            for row in rows:
                content.append(",".join(str(row.get(col, "")) for col in columns))
            return "\n".join(content).encode("utf-8")

    def to_excel(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        ws["A1"] = self.title
        ws["A1"].font = Font(size=16, bold=True)
        ws.merge_cells("A1:F1")

        ws["A2"] = f"Company: {self.company_name}"
        ws["A3"] = f"Generated: {self.generated_at.strftime('%Y-%m-%d %H:%M:%S')}"

        columns = self.data.get("columns", [])
        rows = self.data.get("data", [])

        header_row = 5
        for idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=header_row, column=idx, value=col_name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for row_idx, row in enumerate(rows, start=header_row + 1):
            for col_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

        for col_idx, col_name in enumerate(columns, start=1):
            max_length = len(str(col_name))
            for row_idx in range(header_row + 1, header_row + 1 + len(rows)):
                cell_val = ws.cell(row=row_idx, column=col_idx).value
                cell_text = "" if cell_val is None else str(cell_val)
                if len(cell_text) > max_length:
                    max_length = len(cell_text)
            column_letter = get_column_letter(col_idx)
            ws.column_dimensions[column_letter].width = min(max_length + 2, 60)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def to_csv(self) -> str:
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["Report", self.title])
        writer.writerow(["Company", self.company_name])
        writer.writerow(["Generated", self.generated_at.strftime("%Y-%m-%d %H:%M:%S")])
        writer.writerow([])

        columns = self.data.get("columns", [])
        rows = self.data.get("data", [])
        writer.writerow(columns)
        for row in rows:
            writer.writerow([row.get(col, "") for col in columns])

        return output.getvalue()

    def to_html(self) -> str:
        columns = self.data.get("columns", [])
        rows = self.data.get("data", [])

        head_cells = "".join(f"<th>{escape(str(col))}</th>" for col in columns)
        body_rows = ""
        for row in rows:
            cells = "".join(f"<td>{escape(str(row.get(col, '')))}</td>" for col in columns)
            body_rows += f"<tr>{cells}</tr>"

        return f"""
<!DOCTYPE html>
<html>
<head>
  <meta charset=\"utf-8\" />
  <title>{escape(self.title)}</title>
  <style>
    body {{ font-family: Segoe UI, sans-serif; margin: 24px; color: #152238; }}
    h1 {{ margin-bottom: 4px; }}
    .meta {{ color: #4f5d75; margin-bottom: 16px; }}
    table {{ border-collapse: collapse; width: 100%; }}
    th, td {{ border: 1px solid #d5ddea; padding: 8px; text-align: left; }}
    th {{ background: #203a5f; color: #fff; }}
    tr:nth-child(even) {{ background: #f4f7fb; }}
  </style>
</head>
<body>
  <h1>{escape(self.title)}</h1>
  <div class=\"meta\">Company: {escape(self.company_name)}<br/>Generated: {self.generated_at.strftime('%Y-%m-%d %H:%M:%S')}</div>
  <table>
    <thead><tr>{head_cells}</tr></thead>
    <tbody>{body_rows}</tbody>
  </table>
</body>
</html>
""".strip()
